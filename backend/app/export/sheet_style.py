"""Formatting shared by the sheets of the exported workbook."""
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
WRAP = Alignment(wrap_text=True, vertical="top")


def style(sheet, header_row: int, widths: list[int]) -> None:
    """Column widths, wrapped cells from the header down, a bold header row and the
    rows above it frozen."""
    for n, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(n)].width = width
    for row in sheet.iter_rows(min_row=header_row):
        for cell in row:
            cell.alignment = WRAP
    bold_row(sheet, header_row)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)


def bold_row(sheet, row: int) -> None:
    for cell in sheet[row]:
        cell.font = BOLD


def pages(item: dict) -> str:
    start, end = item["from_page"], item["to_page"]
    return str(start) if start == end else f"{start}–{end}"


def page_ranges(numbers: list[int]) -> str:
    """Pages with each run of consecutive pages as a range: [30, 32, 33, 34] -> "30, 32–34"."""
    runs: list[list[int]] = []
    for n in sorted(set(numbers)):
        if runs and n == runs[-1][1] + 1:
            runs[-1][1] = n
        else:
            runs.append([n, n])
    return ", ".join(str(a) if a == b else f"{a}–{b}" for a, b in runs)
