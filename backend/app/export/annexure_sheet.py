"""The committee's evaluation sheet as an Excel workbook (openpyxl).

Summary: one row per criterion in the RFP's order (group headings with the sum of
their sub-criteria; committee-scored criteria such as a presentation with the
committee's marks), one column per participant in rank order, then the totals and
rank. After it, one sheet per participant in rank order (firm_sheet.py) with its
marks, items and full decision log. Built only from the project's data; no RFP is
named in code.
"""
import io
import re
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook

from app.config import TIMEZONE
from app.eligibility import number
from app.export.firm_sheet import firm_sheet
from app.export.sheet_data import criterion_marks, eligibility_word
from app.export.sheet_style import BOLD, TITLE, bold_row, style

SUMMARY = "Summary"
ELIGIBLE = {"qualified": "Qualified", "not_qualified": "Not qualified"}


def build_workbook(data: dict) -> bytes:
    book = Workbook()
    _summary(book.active, data)
    for firm in data["firms"]:
        firm_sheet(book, firm, data)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def file_name(project: dict) -> str:
    safe = re.sub(r"[^A-Za-z0-9 ._-]+", "", project["name"]).strip() or "project"
    return f"{safe} - evaluation sheet.xlsx"


def _summary(sheet, data: dict) -> None:
    project, results = data["project"], data["results"]
    rows = results["rows"]
    sheet.title = SUMMARY
    now = datetime.now(ZoneInfo(TIMEZONE)).strftime("%d %b %Y %H:%M")
    sheet.append([project["name"]])
    sheet.append([" · ".join(v for v in (project["gem_bid_no"], project["department"],
                                          f"bids closed {project['due']}") if v)])
    sheet.append([f"Exported {now} IST. Marks are the committee's final marks; each "
                  "participant's sheet shows its items and decisions."])
    sheet.append([])
    sheet.append(["Code", "Criterion", "Max marks", *[r["name"] for r in rows]])
    _eligibility(sheet, data)
    for c in data["criteria"]:
        sheet.append([c["code"], c["title"], c["max_marks"],
                      *[criterion_marks(c, r, data["criteria"]) for r in rows]])
    sheet.append(["", "Document marks total", results["docs_max"], *[r["docs"] for r in rows]])
    sheet.append(["", "Total", results["docs_max"] + results["committee_max"],
                  *[r["total"] for r in rows]])
    sheet.append(["", "Rank", None, *[r["rank"] for r in rows]])
    style(sheet, header_row=5, widths=[8, 60, 10, *[14] * len(rows)])
    for row in sheet.iter_rows(min_row=sheet.max_row - 2):
        for cell in row:
            cell.font = BOLD
    sheet["A1"].font = TITLE


def _eligibility(sheet, data: dict) -> None:
    """One row per eligibility criterion (the committee's decision for each firm), then
    each firm's eligibility; then one row per required document. Nothing when the tender
    screens nothing."""
    checks = {(f["row"]["submission_id"], c["criterion_id"]): c
              for f in data["firms"] for c in f["eligibility"]}
    rows = data["results"]["rows"]
    for stage, kind in (("ELIGIBILITY", "pass/fail"), ("DOCUMENT", "required")):
        mine = [r for r in data["requirements"] if r["stage"] == stage]
        for req in mine:
            sheet.append([number(req), req["title"], kind,
                          *[eligibility_word(checks.get((r["submission_id"],
                                                         req["criterion_id"])))
                            for r in rows]])
        if mine and stage == "ELIGIBILITY":
            sheet.append(["", "Eligibility", None,
                          *[ELIGIBLE.get(r.get("eligibility"), "pending") for r in rows]])
            bold_row(sheet, sheet.max_row)
