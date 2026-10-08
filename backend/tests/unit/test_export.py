import io
from decimal import Decimal

from openpyxl import load_workbook

from app.export.annexure_sheet import build_workbook, file_name
from app.results import export_blockers

D = Decimal


CODES = [{"code": "A.1", "max_marks": D(5)}]
COMMITTEE = [{"criterion_id": "c", "code": "C", "title": "Presentation", "max_marks": D(35)}]


def row(name, total, presentation=D(30), cells=None):
    return {"name": name, "total": total, "manual": {"c": presentation},
            "cells": {"A.1": {}} if cells is None else cells}


def test_export_waits_for_evaluations_decisions_and_committee_marks():
    assert export_blockers([], 0, CODES, COMMITTEE, 0) == [
        "No participant has been evaluated yet."]
    assert export_blockers([row("EY", D(90)), row("GT", None)], 2, CODES, COMMITTEE, 0) == [
        "No finished evaluation yet for GT.", "2 marks are not approved by the committee yet."]
    assert export_blockers([row("EY", D(60), presentation=None)], 0, CODES, COMMITTEE, 0) == [
        "C marks are not entered for EY."]
    assert export_blockers([row("EY", D(60), presentation=None)], 0, CODES, [], 0) == []


def test_export_waits_for_eligibility_decisions_but_not_for_firms_not_qualified():
    out = dict(row("Firm C", None), eligibility="not_qualified")      # never evaluated
    pending = dict(row("Firm D", None), eligibility="open")
    assert export_blockers([row("Firm A", D(80)), out], 0, CODES, COMMITTEE, 0) == []
    assert export_blockers([row("Firm A", D(80)), out, pending], 0, CODES, COMMITTEE, 3) == [
        "3 eligibility checks are not decided by the committee yet.",
        "No finished evaluation yet for Firm D."]


def test_a_criterion_with_marks_that_was_never_scored_blocks_the_export():
    assert export_blockers([row("EY", D(60), cells={})], 0, CODES, COMMITTEE, 0) == [
        "A.1 was not scored for EY (evaluated before it could be): evaluate again."]


CRITERIA = [
    {"criterion_id": "a", "code": "A", "parent_code": None, "stage": "TECHNICAL",
     "title": "Experience", "max_marks": D(20)},
    {"criterion_id": "a1", "code": "A.1", "parent_code": "A", "stage": "TECHNICAL",
     "title": "Projects", "max_marks": D(12)},
    {"criterion_id": "a2", "code": "A.2", "parent_code": "A", "stage": "TECHNICAL",
     "title": "Large projects", "max_marks": D(12)},
    {"criterion_id": "c", "code": "C", "parent_code": None, "stage": "TECHNICAL",
     "title": "Presentation", "max_marks": D(10)}]      # committee-scored, any stage
PROJECT = {"name": "Stadium / PMU 2026", "gem_bid_no": "GEM/1", "department": "Dept",
           "due": "07 May 2026"}


def _firm() -> dict:
    cell = {"score_id": "s1", "needs_review": False, "reviewed": True}
    row = {"submission_id": "b1", "name": "Firm/One", "legal_name": "Firm One Pvt Ltd",
           "rank": "1", "manual": {"c": D(8)}, "docs": D("22.5"), "total": D("30.5"),
           "eligibility": "qualified",
           "cells": {"A.1": dict(cell, marks=D(12)), "A.2": dict(cell, marks=D("10.50"))}}
    items = [{"code": "A.1", "title": "Stadium PMU", "label": "A.1 p.10-20", "from_page": 10,
              "to_page": 20, "counted": True, "marks": D(2), "reason": "Completed PMU",
              "action": "OVERRIDE", "decision_reason": "Certificate accepted on review",
              "final_counted": True, "final_marks": D(2),
              "document_notes": ["GSTIN 29ABCDE1234F1Z5: The check digit does not match."]}]
    decisions = [{"code": "A.1", "action": "OVERRIDE", "final_marks": D(2),
                  "reason": "Certificate accepted on review", "full_name": "Local user",
                  "decided": "29 Sep 2026 12:00", "title": "Stadium PMU",
                  "label": "A.1 p.10-20", "from_page": 10, "to_page": 20}]
    checks = [{"criterion_id": "e1", "stage": "ELIGIBILITY", "result": "UNSURE",
               "decision": "MET", "pages": [44, 41, 42, 43, 46],
               "finding": "Two of three years are on the CA certificate.",
               "decision_reason": "Third year verified on the audited statement"},
              {"criterion_id": "e2", "stage": "ELIGIBILITY", "result": "MET",
               "decision": "NOT_MET", "pages": [7], "finding": "Bid form on p.7.",
               "decision_reason": "Only an unsigned copy is in the bid",
               "verification": [{"fact": "document identifier", "value_matches": False,
                                 "note": "MISSING_UDIN: A CA's certificate with no UDIN."},
                                {"fact": "awarded_on", "value_matches": False,
                                 "note": "Not a document check."}]}]
    log = [{"code": "E.1", "stage": "ELIGIBILITY", "decision": "MET",
            "full_name": "Local user", "decided": "29 Sep 2026 11:00",
            "reason": "Third year verified"},
           {"code": "E.2", "stage": "ELIGIBILITY", "decision": "NOT_MET",
            "full_name": "Local user", "decided": "29 Sep 2026 11:05", "reason": "Unsigned"}]
    return {"row": row, "items": items, "decisions": decisions, "eligibility": checks,
            "eligibility_log": log}


def _book():
    firm = _firm()
    results = {"committee_max": D(10), "docs_max": D(24), "rows": [firm["row"]]}
    return load_workbook(io.BytesIO(build_workbook(
        {"project": PROJECT, "results": results, "criteria": CRITERIA, "firms": [firm],
         "requirements": [
             {"criterion_id": "e1", "code": "E.1", "number": "1", "stage": "ELIGIBILITY",
              "title": "Turnover"},
             {"criterion_id": "e2", "code": "E.2", "number": "2", "stage": "ELIGIBILITY",
              "title": "Bid form"}],
         "history": [], "mark_history": [], "events": [], "approvals": []})))


def test_the_summary_compares_firms_on_eligibility_and_marks():
    book = _book()
    assert book.sheetnames == ["Summary", "Firm One", "Audit log"]   # "/" not allowed
    rows = {r[1]: r for r in book["Summary"].iter_rows(min_row=6, values_only=True) if r[1]}
    assert rows["Turnover"][2:4] == ("pass/fail", "met")
    assert rows["Eligibility"][3] == "Qualified"
    assert rows["Bid form"][2:4] == ("pass/fail", "not met")
    assert rows["Experience"][3] == D("22.5") and rows["Projects"][3] == 12
    assert rows["Large projects"][3] == 10.5 and rows["Presentation"][3] == 8
    assert rows["Total"][2:4] == (34, 30.5)
    assert file_name(PROJECT) == "Stadium  PMU 2026 - evaluation sheet.xlsx"


def test_each_firm_sheet_has_its_eligibility_items_and_full_decision_log():
    firm = list(_book()["Firm One"].iter_rows(values_only=True))
    assert firm[0][0] == "Firm One Pvt Ltd (Firm/One)" and "total 30.5 of 34" in firm[1][0]
    assert firm[4][:7] == ("1", "Turnover", "41–44, 46", None, None, "unsure", "met")
    assert "Committee: Third year verified" in firm[4][8]
    assert firm[5][:7] == ("2", "Bid form", "7", None, None, "met", "not met")
    # A check's document problems are on its own row, between the finding and the reason.
    assert firm[5][8] == ("Bid form on p.7.\nDocument check: MISSING_UDIN: A CA's certificate "
                          "with no UDIN.\nCommittee: Only an unsigned copy is in the bid")
    assert firm[9][:8] == ("A.1", "Projects", None, 12, None, "approved", None, 12)
    assert firm[10][1:8] == ("Stadium PMU", "10–20", None, 2, "overridden", "yes", 2)
    assert firm[10][8] == ("Completed PMU\nDocument check: GSTIN 29ABCDE1234F1Z5: The check "
                           "digit does not match.\nCommittee: Certificate accepted on review")
    assert firm[-3][:2] == ("1", "eligibility") and firm[-3][5] == "met"
    assert firm[-2][:2] == ("2", "eligibility") and firm[-2][5] == "not met"
    assert firm[-1][:3] == ("A.1", "Stadium PMU", "10–20") and firm[-1][5] == "override"
